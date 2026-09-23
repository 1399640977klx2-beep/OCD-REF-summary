"""Tab: Match作图 - generate standard and custom Match charts."""
import os
import traceback

from PyQt5.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QLineEdit, QPushButton,
    QComboBox, QFileDialog, QGroupBox, QMessageBox, QTextEdit, QSpinBox,
    QTableWidget, QTableWidgetItem, QAbstractItemView, QListWidget,
    QListWidgetItem, QCheckBox,
)
from PyQt5.QtCore import Qt
from openpyxl.utils import get_column_letter

from utils.match_chart_builder import (
    analyze_match_workbook,
    generate_match_charts,
    list_sheets,
    DEFAULT_END_ROW,
)


class MatchChartsTab(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.main_window = parent
        self.analysis = None
        self._last_auto_save = ""
        self._init_ui()

    def _init_ui(self):
        layout = QVBoxLayout(self)

        source_group = QGroupBox("数据源")
        source_layout = QVBoxLayout()

        row1 = QHBoxLayout()
        row1.addWidget(QLabel("Excel 文件:"))
        self.input_edit = QLineEdit()
        self.input_edit.setPlaceholderText("选择需要生成 Match 图表的 Excel")
        row1.addWidget(self.input_edit, 1)
        self.input_btn = QPushButton("浏览")
        self.input_btn.clicked.connect(self._browse_input)
        row1.addWidget(self.input_btn)
        source_layout.addLayout(row1)

        row2 = QHBoxLayout()
        row2.addWidget(QLabel("Sheet:"))
        self.sheet_combo = QComboBox()
        self.sheet_combo.setEnabled(False)
        self.sheet_combo.currentTextChanged.connect(self._analyze)
        row2.addWidget(self.sheet_combo, 1)
        row2.addWidget(QLabel("绘图到第"))
        self.end_row_spin = QSpinBox()
        self.end_row_spin.setRange(1, 1048576)
        self.end_row_spin.setValue(DEFAULT_END_ROW)
        self.end_row_spin.setSingleStep(100)
        row2.addWidget(self.end_row_spin)
        row2.addWidget(QLabel("行"))
        source_layout.addLayout(row2)

        source_group.setLayout(source_layout)
        layout.addWidget(source_group)

        bias_group = QGroupBox("Bias 图选择")
        bias_layout = QVBoxLayout()
        self.bias_info = QLabel("选择文件和 Sheet 后自动识别 Bias 类型")
        self.bias_info.setObjectName("matchChartsBiasInfo")
        bias_layout.addWidget(self.bias_info)
        self.bias_table = QTableWidget()
        self.bias_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.bias_table.setSelectionMode(QAbstractItemView.NoSelection)
        self.bias_table.setMaximumHeight(150)
        bias_layout.addWidget(self.bias_table)
        bias_btns = QHBoxLayout()
        self.bias_all_btn = QPushButton("全选")
        self.bias_all_btn.clicked.connect(lambda: self._set_all_bias(True))
        bias_btns.addWidget(self.bias_all_btn)
        self.bias_none_btn = QPushButton("全不选")
        self.bias_none_btn.clicked.connect(lambda: self._set_all_bias(False))
        bias_btns.addWidget(self.bias_none_btn)
        bias_btns.addStretch()
        bias_layout.addLayout(bias_btns)
        bias_group.setLayout(bias_layout)
        layout.addWidget(bias_group)

        custom_group = QGroupBox("自定义图表")
        custom_layout = QVBoxLayout()
        custom_layout.addWidget(QLabel("全部表头（折线图可多选，作为系列）："))
        self.headers_list = QListWidget()
        self.headers_list.setSelectionMode(
            QAbstractItemView.ExtendedSelection)
        self.headers_list.setMaximumHeight(120)
        custom_layout.addWidget(self.headers_list)

        label_row = QHBoxLayout()
        label_row.addWidget(QLabel("折线图标签列:"))
        self.label_combo = QComboBox()
        label_row.addWidget(self.label_combo, 1)
        custom_layout.addLayout(label_row)

        scatter_row = QHBoxLayout()
        self.custom_scatter_cb = QCheckBox("自定义散点图")
        scatter_row.addWidget(self.custom_scatter_cb)
        scatter_row.addWidget(QLabel("X:"))
        self.scatter_x_combo = QComboBox()
        scatter_row.addWidget(self.scatter_x_combo, 1)
        scatter_row.addWidget(QLabel("Y:"))
        self.scatter_y_combo = QComboBox()
        scatter_row.addWidget(self.scatter_y_combo, 1)
        custom_layout.addLayout(scatter_row)

        self.custom_line_cb = QCheckBox(
            "自定义折线图（使用上方选中的表头作为系列）")
        custom_layout.addWidget(self.custom_line_cb)
        custom_group.setLayout(custom_layout)
        layout.addWidget(custom_group)

        output_group = QGroupBox("输出")
        output_layout = QVBoxLayout()
        save_row = QHBoxLayout()
        save_row.addWidget(QLabel("保存路径:"))
        self.save_edit = QLineEdit()
        self.save_edit.setPlaceholderText("选择输出 xlsx 文件")
        save_row.addWidget(self.save_edit, 1)
        self.save_btn = QPushButton("浏览")
        self.save_btn.clicked.connect(self._browse_save)
        save_row.addWidget(self.save_btn)
        output_layout.addLayout(save_row)

        gen_row = QHBoxLayout()
        self.generate_btn = QPushButton("生成 Match 图表")
        self.generate_btn.setObjectName("btnGenerateMatchCharts")
        self.generate_btn.clicked.connect(self._generate)
        gen_row.addWidget(self.generate_btn)
        gen_row.addStretch()
        output_layout.addLayout(gen_row)
        output_group.setLayout(output_layout)
        layout.addWidget(output_group)

        self.log = QTextEdit()
        self.log.setReadOnly(True)
        self.log.setMaximumHeight(160)
        self.log.setPlaceholderText("分析和生成日志将显示在这里...")
        layout.addWidget(self.log, 1)
        self.setLayout(layout)

    def _browse_input(self):
        fp, _ = QFileDialog.getOpenFileName(
            self, "选择 Excel 文件", self.main_window.last_directory,
            "Excel 文件 (*.xlsx)")
        if not fp:
            return
        self.main_window.last_directory = os.path.dirname(fp)
        self.input_edit.setText(fp)
        self._load_sheets(fp)

    def _load_sheets(self, filepath):
        try:
            sheets = list_sheets(filepath)
        except Exception as e:
            QMessageBox.critical(self, "读取失败", str(e))
            return
        self.sheet_combo.blockSignals(True)
        self.sheet_combo.clear()
        self.sheet_combo.addItems(sheets)
        self.sheet_combo.setEnabled(bool(sheets))
        if sheets:
            preferred = next(
                (name for name in sheets if name.strip().lower() == "match"),
                0,
            )
            if isinstance(preferred, str):
                self.sheet_combo.setCurrentText(preferred)
            else:
                self.sheet_combo.setCurrentIndex(0)
        self.sheet_combo.blockSignals(False)
        if sheets:
            self._analyze()

    def _analyze(self):
        input_path = self.input_edit.text().strip()
        sheet = self.sheet_combo.currentText().strip()
        if not input_path or not sheet:
            return
        try:
            self.analysis = analyze_match_workbook(
                input_path, sheet, self.end_row_spin.value())
        except Exception as e:
            self.analysis = None
            self.log.append("分析失败: " + str(e))
            return

        self._populate_bias_table()
        self._populate_custom_choices()
        self._update_default_save()

        self.log.clear()
        self.log.append(
            f"识别完成: 表头第 {self.analysis.header_row} 行，"
            f"数据 {self.analysis.data_start_row}:{self.analysis.end_row}")
        for w in self.analysis.warnings:
            self.log.append("[警告] " + w)
        if not self.analysis.warnings:
            self.log.append("未发现 Bias 类型缺失")

    def _populate_bias_table(self):
        params = self.analysis.parameters
        variants = []
        for param in params:
            for variant in param.bias_cols:
                if variant not in variants:
                    variants.append(variant)

        self.bias_table.clear()
        self.bias_table.setRowCount(len(params))
        self.bias_table.setColumnCount(len(variants))
        self.bias_table.setVerticalHeaderLabels([p.name for p in params])
        self.bias_table.setHorizontalHeaderLabels(variants)
        for row, param in enumerate(params):
            for col, variant in enumerate(variants):
                if variant in param.bias_cols:
                    item = QTableWidgetItem()
                    item.setFlags(Qt.ItemIsEnabled | Qt.ItemIsUserCheckable)
                    item.setCheckState(Qt.Checked)
                else:
                    item = QTableWidgetItem("—")
                    item.setFlags(Qt.NoItemFlags)
                self.bias_table.setItem(row, col, item)
        self.bias_table.resizeColumnsToContents()
        self.bias_info.setText(
            f"识别到 {len(params)} 个参数，Bias 类型: "
            f"{', '.join(variants) if variants else '无'}")

    def _populate_custom_choices(self):
        headers = [
            (col, name) for col, name in self.analysis.headers.items() if name
        ]
        self.headers_list.clear()
        for col, name in headers:
            item = QListWidgetItem(f"{get_column_letter(col)}: {name}")
            item.setData(Qt.UserRole, col)
            self.headers_list.addItem(item)

        self.label_combo.clear()
        self.label_combo.addItem("J:M 多级标签", None)
        for col, name in headers:
            self.label_combo.addItem(f"{get_column_letter(col)}: {name}", col)

        self.scatter_x_combo.clear()
        self.scatter_y_combo.clear()
        for col, name in headers:
            self.scatter_x_combo.addItem(f"{get_column_letter(col)}: {name}", col)
            self.scatter_y_combo.addItem(f"{get_column_letter(col)}: {name}", col)
        if self.analysis.parameters:
            p = self.analysis.parameters[0]
            if p.raw_col is not None:
                self._select_combo_data(self.scatter_x_combo, p.raw_col)
            if p.ref_col is not None:
                self._select_combo_data(self.scatter_y_combo, p.ref_col)

    @staticmethod
    def _select_combo_data(combo, value):
        idx = combo.findData(value)
        if idx >= 0:
            combo.setCurrentIndex(idx)

    def _set_all_bias(self, checked):
        state = Qt.Checked if checked else Qt.Unchecked
        for row in range(self.bias_table.rowCount()):
            for col in range(self.bias_table.columnCount()):
                item = self.bias_table.item(row, col)
                if item and item.flags() & Qt.ItemIsUserCheckable:
                    item.setCheckState(state)

    def _collect_bias_selection(self):
        selected = {}
        for row, param in enumerate(self.analysis.parameters):
            variants = []
            for col in range(self.bias_table.columnCount()):
                item = self.bias_table.item(row, col)
                if (item and item.flags() & Qt.ItemIsUserCheckable
                        and item.checkState() == Qt.Checked):
                    variants.append(self.bias_table.horizontalHeaderItem(col).text())
            selected[param.name] = variants
        return selected

    def _collect_custom_charts(self):
        custom = {"scatter": [], "line": []}
        if self.custom_scatter_cb.isChecked():
            x_col = self.scatter_x_combo.currentData()
            y_col = self.scatter_y_combo.currentData()
            if x_col is not None and y_col is not None:
                custom["scatter"].append({"x_col": x_col, "y_col": y_col})

        if self.custom_line_cb.isChecked():
            columns = [
                item.data(Qt.UserRole)
                for item in self.headers_list.selectedItems()
            ]
            if not columns:
                raise ValueError("自定义折线图未选择任何表头")
            custom["line"].append({
                "columns": columns,
                "label_col": self.label_combo.currentData(),
                "title": "Custom_Line",
            })
        return custom

    def _suggest_save_path(self):
        input_path = self.input_edit.text().strip()
        if not input_path:
            return ""
        stem = os.path.splitext(os.path.basename(input_path))[0]
        return os.path.join(
            os.path.dirname(input_path), stem + "_match_charts.xlsx")

    def _update_default_save(self):
        suggested = self._suggest_save_path()
        if not suggested:
            return
        current = self.save_edit.text().strip()
        if current == "" or current == self._last_auto_save:
            self.save_edit.setText(suggested)
            self._last_auto_save = suggested

    def _browse_save(self):
        current = self.save_edit.text().strip()
        default_path = current or self._suggest_save_path()
        if not default_path:
            default_path = self.main_window.last_directory
        fp, _ = QFileDialog.getSaveFileName(
            self, "选择保存路径", default_path, "Excel 文件 (*.xlsx)")
        if not fp:
            return
        if not fp.lower().endswith(".xlsx"):
            fp += ".xlsx"
        self.save_edit.setText(fp)
        self._last_auto_save = fp

    def _generate(self):
        input_path = self.input_edit.text().strip()
        sheet = self.sheet_combo.currentText().strip()
        save_path = self.save_edit.text().strip()
        if not input_path or not os.path.isfile(input_path):
            QMessageBox.warning(self, "错误", "请选择有效的 Excel 文件")
            return
        if not sheet:
            QMessageBox.warning(self, "错误", "请选择 Sheet")
            return
        if not save_path:
            QMessageBox.warning(self, "错误", "请选择保存路径")
            return
        if not save_path.lower().endswith(".xlsx"):
            save_path += ".xlsx"
        if self.analysis is None:
            self._analyze()
        if self.analysis is None:
            return

        try:
            bias_selection = self._collect_bias_selection()
            custom_charts = self._collect_custom_charts()
        except Exception as e:
            QMessageBox.warning(self, "错误", str(e))
            return

        self.log.clear()
        self.log.append(f"输入: {input_path} / {sheet}")
        self.log.append(f"绘图行范围: {self.analysis.data_start_row}:"
                        f"{self.analysis.end_row}")
        self.main_window.set_status("正在生成 Match 图表...")
        try:
            result = generate_match_charts(
                input_path, save_path, sheet,
                self.end_row_spin.value(),
                bias_selection=bias_selection,
                custom_charts=custom_charts)
            for warning in result["warnings"]:
                self.log.append("[警告] " + warning)
            self.log.append(
                f"完成: 标准图表 {result['standard_charts']} 张，"
                f"自定义图表 {result['custom_charts']} 张")
            self.log.append("已保存: " + result["output_path"])
            self.main_window.set_status(
                "Match 图表已生成: " + os.path.basename(result["output_path"]))
            QMessageBox.information(
                self, "完成", f"Match 图表已生成：\n{result['output_path']}")
        except Exception as e:
            self.log.append("错误: " + str(e))
            self.log.append(traceback.format_exc())
            self.main_window.set_status("Match 图表生成失败")
            QMessageBox.critical(self, "失败", str(e))
